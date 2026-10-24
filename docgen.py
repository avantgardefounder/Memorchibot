"""Генерация Word-документов (ТЗ, акты, письма, договоры, пояснительные записки, сметы)."""
import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

OUT_DIR = Path(__file__).with_name("output")
OUT_DIR.mkdir(exist_ok=True)

COMPANY_NAME = ""  # можно задать через .env -> COMPANY_NAME


def _safe(name):
    name = re.sub(r"[^\w\-а-яА-ЯёЁʻʼ' ]+", "", name).strip().replace(" ", "_")
    return name[:60] or "document"


def create_document(title, sections, doc_type="other", table=None, signatures=None,
                    city="Ташкент", company=None):
    """sections: [{heading, text}] — text может содержать переносы строк и пункты '- '.
    table: {columns: [...], rows: [[...], ...]} — необязательно.
    signatures: ["Заказчик", "Подрядчик"] — блоки подписей внизу."""
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for s in d.sections:
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.top_margin = s.bottom_margin = Cm(2)

    comp = company or COMPANY_NAME
    if comp:
        p = d.add_paragraph(comp)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.runs[0].italic = True

    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(14)

    meta = d.add_paragraph()
    meta.add_run(f"г. {city}").italic = True
    meta.add_run("\t\t\t\t\t\t")
    meta.add_run(dt.date.today().strftime("%d.%m.%Y")).italic = True

    for sec in sections:
        if sec.get("heading"):
            hp = d.add_paragraph()
            hr = hp.add_run(sec["heading"])
            hr.bold = True
        for line in str(sec.get("text", "")).split("\n"):
            line = line.rstrip()
            if not line:
                continue
            if line.lstrip().startswith(("- ", "• ")):
                d.add_paragraph(line.lstrip()[2:], style="List Bullet")
            else:
                para = d.add_paragraph(line)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if table and table.get("columns"):
        cols = table["columns"]
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]
            cell.text = str(c)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in table.get("rows", []):
            cells = t.add_row().cells
            for i, v in enumerate(row[:len(cols)]):
                cells[i].text = str(v)

    if signatures:
        d.add_paragraph()
        for who in signatures:
            d.add_paragraph(f"{who}: ____________________ / ____________________ /")
            d.add_paragraph("М.П.")

    fname = OUT_DIR / f"{_safe(title)}_{dt.datetime.now():%Y%m%d_%H%M%S}.docx"
    d.save(fname)
    return str(fname)


def estimate_to_document(est, title="Смета"):
    cols = ["№", "Наименование", "Ед.", "Кол-во", f"Цена, {est['currency']}", f"Сумма, {est['currency']}"]
    rows = [[i + 1, l["name"], l["unit"], l["qty"], f"{l['price']:,}".replace(",", " "),
             f"{l['sum']:,}".replace(",", " ")] for i, l in enumerate(est["lines"])]
    rows.append(["", "Итого", "", "", "", f"{est['subtotal']:,}".replace(",", " ")])
    rows.append(["", f"Непредвиденные расходы {est['contingency_percent']}%", "", "", "",
                 f"{est['contingency']:,}".replace(",", " ")])
    rows.append(["", "ВСЕГО", "", "", "", f"{est['total']:,}".replace(",", " ")])
    note = (f"Регион: {est['region']} (коэф. {est['region_coef']}). Прайс от {est['price_list_date']}. "
            "Смета ориентировочная, не является проектно-сметной документацией.")
    return create_document(title, [{"heading": "", "text": note}], doc_type="estimate",
                           table={"columns": cols, "rows": rows})
