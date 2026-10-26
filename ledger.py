"""Учёт средств клиента по объектам и отчёты для клиента (Word + Excel).

Модель: фирма строит за счёт клиента. Клиент вносит деньги (поступления),
фирма тратит их на материалы, работы, технику и т.д. (расходы) и, если договорились,
получает вознаграждение — процент от расходов.
Остаток средств клиента = поступления − расходы − вознаграждение.
Все суммы считаются здесь, в коде, а не AI.
"""
import datetime as dt
import re
import time
from collections import OrderedDict
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import db
from docgen import OUT_DIR
from translit import to_cyr

CURRENCIES = ("UZS", "USD")

TZ = ZoneInfo("Asia/Tashkent")
EXPENSE_CATEGORIES = ["materials", "works", "equipment", "transport", "design", "permits", "other"]
INCOME_CATEGORIES = ["advance", "payment", "other_income"]

L = {  # подписи в отчёте
    "ru": {
        "title": "ОТЧЁТ ОБ ИСПОЛЬЗОВАНИИ СРЕДСТВ ЗАКАЗЧИКА",
        "object": "Объект", "client": "Заказчик", "address": "Адрес", "period": "Период",
        "all_time": "за весь период", "report_date": "Дата отчёта", "contractor": "Подрядчик",
        "summary": "1. Сводка", "budget": "Бюджет объекта", "received": "Получено от заказчика",
        "spent": "Израсходовано", "fee": "Вознаграждение подрядчика", "balance": "Остаток средств заказчика",
        "overspend": "Перерасход (к доплате заказчиком)", "budget_used": "Освоено бюджета",
        "by_cat": "2. Расходы по статьям", "category": "Статья", "amount": "Сумма", "share": "Доля",
        "incomes": "3. Поступления от заказчика", "expenses": "4. Детализация расходов",
        "date": "Дата", "desc": "Наименование", "qty": "Кол-во", "unit": "Ед.", "price": "Цена",
        "type": "Тип", "total": "Итого", "no": "№",
        "note": "Отчёт сформирован по данным учёта подрядчика. Подтверждающие документы (чеки, накладные, акты) предоставляются по запросу.",
        "sign_c": "Подрядчик", "sign_z": "Заказчик", "of_expenses": "от расходов",
        "cats": {"materials": "Материалы", "works": "Работы", "equipment": "Техника и инструмент",
                 "transport": "Доставка и транспорт", "design": "Проектирование", "permits": "Разрешения и согласования",
                 "other": "Прочее", "advance": "Аванс", "payment": "Оплата", "other_income": "Прочее поступление"},
        "sheet_sum": "Сводка", "sheet_inc": "Поступления", "sheet_exp": "Расходы",
    },
    "uz": {
        "title": "BUYURTMACHI MABLAG'LARIDAN FOYDALANISH TO'G'RISIDA HISOBOT",
        "object": "Obyekt", "client": "Buyurtmachi", "address": "Manzil", "period": "Davr",
        "all_time": "butun davr uchun", "report_date": "Hisobot sanasi", "contractor": "Pudratchi",
        "summary": "1. Umumiy ma'lumot", "budget": "Obyekt byudjeti", "received": "Buyurtmachidan olingan",
        "spent": "Sarflangan", "fee": "Pudratchi mukofoti", "balance": "Buyurtmachi mablag'lari qoldig'i",
        "overspend": "Ortiqcha sarf (buyurtmachi to'lashi kerak)", "budget_used": "Byudjetning o'zlashtirilgani",
        "by_cat": "2. Moddalar bo'yicha xarajatlar", "category": "Modda", "amount": "Summa", "share": "Ulush",
        "incomes": "3. Buyurtmachidan tushumlar", "expenses": "4. Xarajatlar tafsiloti",
        "date": "Sana", "desc": "Nomi", "qty": "Miqdori", "unit": "O'lch.", "price": "Narxi",
        "type": "Turi", "total": "Jami", "no": "№",
        "note": "Hisobot pudratchi hisobi ma'lumotlari asosida tuzilgan. Tasdiqlovchi hujjatlar (cheklar, yuk xatlari, dalolatnomalar) so'rov bo'yicha taqdim etiladi.",
        "sign_c": "Pudratchi", "sign_z": "Buyurtmachi", "of_expenses": "xarajatlardan",
        "cats": {"materials": "Materiallar", "works": "Ishlar", "equipment": "Texnika va asboblar",
                 "transport": "Yetkazib berish va transport", "design": "Loyihalash", "permits": "Ruxsatnomalar va kelishuvlar",
                 "other": "Boshqa", "advance": "Avans", "payment": "To'lov", "other_income": "Boshqa tushum"},
        "sheet_sum": "Umumiy", "sheet_inc": "Tushumlar", "sheet_exp": "Xarajatlar",
    },
    "en": {
        "title": "CLIENT FUNDS USAGE REPORT",
        "object": "Project", "client": "Client", "address": "Address", "period": "Period",
        "all_time": "entire period", "report_date": "Report date", "contractor": "Contractor",
        "summary": "1. Summary", "budget": "Project budget", "received": "Received from client",
        "spent": "Spent", "fee": "Contractor fee", "balance": "Client funds balance",
        "overspend": "Overspend (payable by client)", "budget_used": "Budget used",
        "by_cat": "2. Expenses by category", "category": "Category", "amount": "Amount", "share": "Share",
        "incomes": "3. Payments from client", "expenses": "4. Expense details",
        "date": "Date", "desc": "Description", "qty": "Qty", "unit": "Unit", "price": "Price",
        "type": "Type", "total": "Total", "no": "#",
        "note": "This report is based on the contractor's records. Supporting documents (receipts, delivery notes, acts) are available on request.",
        "sign_c": "Contractor", "sign_z": "Client", "of_expenses": "of expenses",
        "cats": {"materials": "Materials", "works": "Labour / works", "equipment": "Equipment & tools",
                 "transport": "Delivery & transport", "design": "Design", "permits": "Permits & approvals",
                 "other": "Other", "advance": "Advance", "payment": "Payment", "other_income": "Other income"},
        "sheet_sum": "Summary", "sheet_inc": "Payments", "sheet_exp": "Expenses",
    },
}


L["uzc"] = to_cyr(L["uz"])
PAID_BY = ("client", "foreman", "company", "other")


def _today():
    return dt.datetime.now(TZ).date().isoformat()


def _rnd(x, cur="UZS"):
    return round(x) if cur == "UZS" else round(x, 2)


def _money(x, cur="UZS"):
    if cur == "UZS":
        return f"{round(x):,}".replace(",", " ")
    return f"{x:,.2f}".replace(",", " ")


def _parse_date(s):
    if not s:
        return _today()
    s = s.strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Не понял дату: {s}. Формат ДД.ММ.ГГГГ")


# ---------- объекты ----------
def _obj_row(r):
    keys = ["id", "name", "client", "address", "budget", "fee_percent", "currency", "status", "created",
            "obj_type", "area", "floors", "lat", "lon", "radius", "start_date", "end_date", "group_chat_id"]
    return dict(zip(keys, r))


OBJECT_FIELDS = ("name", "client", "address", "budget", "fee_percent", "currency", "status", "obj_type", "area",
                 "floors", "lat", "lon", "radius", "start_date", "end_date", "group_chat_id")


def create_object(name, client="", address="", budget=0, fee_percent=0, currency="UZS", **extra):
    currency = (currency or "UZS").upper()
    if currency not in CURRENCIES:
        raise ValueError("Валюта объекта может быть только UZS или USD")
    db._con.execute("INSERT INTO objects(name,client,address,budget,fee_percent,currency,created) VALUES(?,?,?,?,?,?,?)",
                    (name.strip(), client, address, float(budget or 0), float(fee_percent or 0), currency, time.time()))
    db._con.commit()
    o = find_object(name.strip())
    extra = {k: v for k, v in extra.items() if k in OBJECT_FIELDS and v not in (None, "")}
    return update_object(o["id"], **extra) if extra else o


def update_object(object, **fields):
    o = find_object(object)
    allowed = {k: v for k, v in fields.items() if k in OBJECT_FIELDS and v is not None}
    for k in ("start_date", "end_date"):
        if allowed.get(k):
            allowed[k] = _parse_date(allowed[k])
    if not allowed:
        return o
    if "currency" in allowed:
        allowed["currency"] = str(allowed["currency"]).upper()
        if allowed["currency"] not in CURRENCIES:
            raise ValueError("Валюта объекта может быть только UZS или USD")
        if allowed["currency"] != o["currency"] and _tx(o["id"]):
            raise ValueError("Нельзя сменить валюту объекта, по которому уже есть операции")
    sets = ",".join(f"{k}=?" for k in allowed)
    db._con.execute(f"UPDATE objects SET {sets} WHERE id=?", (*allowed.values(), o["id"]))
    db._con.commit()
    return find_object(o["id"])


def list_objects(include_closed=False, ids=None, with_finance=True):
    q = "SELECT * FROM objects" + ("" if include_closed else " WHERE status='active'") + " ORDER BY id"
    out = []
    for r in db._con.execute(q).fetchall():
        o = _obj_row(r)
        if ids is not None and o["id"] not in ids:
            continue
        if not with_finance:
            out.append({"id": o["id"], "name": o["name"], "client": o["client"], "status": o["status"],
                        "address": o["address"]})
            continue
        s = summary(o["id"])
        out.append({"id": o["id"], "name": o["name"], "client": o["client"], "status": o["status"], "currency": o["currency"],
                    "received": s["received"], "spent": s["spent"], "balance": s["balance"]})
    return out


def find_object(ref):
    """По id или по названию (без учёта регистра, допускается часть названия)."""
    if isinstance(ref, int) or (isinstance(ref, str) and ref.isdigit()):
        r = db._con.execute("SELECT * FROM objects WHERE id=?", (int(ref),)).fetchone()
        if r:
            return _obj_row(r)
    rows = [_obj_row(r) for r in db._con.execute("SELECT * FROM objects").fetchall()]
    ref_l = str(ref).strip().lower()
    exact = [o for o in rows if o["name"].lower() == ref_l]
    if exact:
        return exact[0]
    part = [o for o in rows if ref_l in o["name"].lower() or ref_l in (o["client"] or "").lower()]
    if len(part) == 1:
        return part[0]
    if len(part) > 1:
        raise ValueError("Подходит несколько объектов: " + ", ".join(o["name"] for o in part) + ". Уточните.")
    raise ValueError(f"Объект «{ref}» не найден. Существующие: " + (", ".join(o["name"] for o in rows) or "нет"))


# ---------- операции ----------
def add_transactions(object, items, author_id=None, author_name="", receipt_file_id=None):
    """items: [{type: income|expense, category, description, date?, amount? | qty+price, unit?}]"""
    o = find_object(object)
    before = summary(o["id"])
    added = []
    for it in items:
        ttype = it.get("type", "expense")
        if ttype not in ("income", "expense"):
            raise ValueError("type: income или expense")
        cat = it.get("category") or ("payment" if ttype == "income" else "other")
        valid = INCOME_CATEGORIES if ttype == "income" else EXPENSE_CATEGORIES
        if cat not in valid:
            cat = "other_income" if ttype == "income" else "other"
        qty, price = it.get("qty"), it.get("price")
        if qty not in (None, "") and price not in (None, ""):
            orig = float(qty) * float(price)
        elif it.get("amount") not in (None, ""):
            orig = float(it["amount"])
            qty = price = None
        else:
            raise ValueError(f"Нет суммы для «{it.get('description', '')}»: нужна amount или qty и price")
        if orig <= 0:
            raise ValueError("Сумма должна быть больше нуля")
        date = _parse_date(it.get("date"))
        ccy = (it.get("currency") or o["currency"]).upper()
        if ccy not in CURRENCIES:
            raise ValueError("Валюта операции может быть только UZS или USD")
        rate = None
        if ccy == o["currency"]:
            amount = orig
            orig_ccy = orig_amount = None
        else:
            if not it.get("rate"):
                raise ValueError("Операция в другой валюте: нужен курс (сколько сумов за 1 USD). Спросите курс у пользователя.")
            rate = float(it["rate"])  # всегда сумов за 1 USD
            if rate <= 0:
                raise ValueError("Курс должен быть больше нуля")
            factor = rate if ccy == "USD" else 1 / rate  # USD→UZS: ×курс, UZS→USD: ÷курс
            amount = orig * factor
            orig_ccy, orig_amount = ccy, orig
            if price not in (None, ""):
                price = float(price) * factor  # цена в валюте объекта, чтобы кол-во × цена = сумма
        paid_by = it.get("paid_by") or None
        if paid_by and paid_by not in PAID_BY:
            paid_by = "other"
        receipt = receipt_file_id if ttype == "expense" else None
        cur = db._con.execute(
            "INSERT INTO transactions(object_id,date,type,category,description,qty,unit,price,amount,author_id,author_name,created,"
            "orig_currency,orig_amount,rate,paid_by,receipt_file_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (o["id"], date, ttype, cat, it.get("description", ""), qty, it.get("unit", ""), price, amount,
             author_id, author_name, time.time(), orig_ccy, orig_amount, rate, paid_by, receipt))
        rec = {"id": cur.lastrowid, "date": date, "type": ttype, "category": cat,
               "description": it.get("description", ""), "amount": _rnd(amount, o["currency"]), "currency": o["currency"]}
        if orig_ccy:
            rec.update({"original": f"{_money(orig_amount, orig_ccy)} {orig_ccy}",
                        "rate_uzs_per_usd": rate})
        if paid_by:
            rec["paid_by"] = paid_by
        if receipt:
            rec["receipt_photo"] = "прикреплено"
        added.append(rec)
    db._con.commit()
    s = summary(o["id"])
    res = {"object": o["name"], "object_id": o["id"], "object_currency": o["currency"], "added": added,
           "balance_after": s["balance"]}
    budget = o["budget"] or 0
    if budget and before["spent"] + before["fee"] <= budget < s["spent"] + s["fee"]:
        res["budget_exceeded"] = {"budget": s["budget"], "used": _rnd(s["spent"] + s["fee"], o["currency"])}
    if s["balance"] < 0 <= before["balance"]:
        res["client_funds_negative"] = s["balance"]
    return res


def tx_object_id(transaction_id):
    r = db._con.execute("SELECT object_id FROM transactions WHERE id=? AND deleted=0", (int(transaction_id),)).fetchone()
    if not r:
        raise ValueError(f"Операция №{transaction_id} не найдена")
    return r[0]


def receipts(object_id, ids=None, date_from=None, date_to=None):
    rows = [r for r in _tx(object_id, date_from, date_to) if r.get("receipt_file_id")]
    if ids:
        ids = {int(i) for i in ids}
        rows = [r for r in rows if r["id"] in ids]
    return rows


def delete_transaction(transaction_id):
    r = db._con.execute("SELECT id, description, amount FROM transactions WHERE id=? AND deleted=0",
                        (int(transaction_id),)).fetchone()
    if not r:
        raise ValueError(f"Операция №{transaction_id} не найдена")
    db._con.execute("UPDATE transactions SET deleted=1 WHERE id=?", (int(transaction_id),))
    db._con.commit()
    return {"deleted": {"id": r[0], "description": r[1], "amount": round(r[2], 2)}}


def _tx(object_id, date_from=None, date_to=None):
    q = ("SELECT id,date,type,category,description,qty,unit,price,amount,author_name,orig_currency,orig_amount,rate,"
         "paid_by,receipt_file_id FROM transactions "
         "WHERE object_id=? AND deleted=0")
    args = [object_id]
    if date_from:
        q += " AND date>=?"
        args.append(_parse_date(date_from))
    if date_to:
        q += " AND date<=?"
        args.append(_parse_date(date_to))
    q += " ORDER BY date, id"
    keys = ["id", "date", "type", "category", "description", "qty", "unit", "price", "amount", "author",
            "orig_currency", "orig_amount", "rate", "paid_by", "receipt_file_id"]
    return [dict(zip(keys, r)) for r in db._con.execute(q, args).fetchall()]


def list_transactions(object, last=20, date_from=None, date_to=None):
    o = find_object(object)
    rows = _tx(o["id"], date_from, date_to)[-int(last):]
    out = []
    for r in rows:
        r = {**r, "amount": _rnd(r["amount"], o["currency"]), "receipt_photo": bool(r.pop("receipt_file_id"))}
        out.append(r)
    return {"object": o["name"], "currency": o["currency"], "transactions": out}


def summary(object, date_from=None, date_to=None):
    o = find_object(object)
    rows = _tx(o["id"], date_from, date_to)
    received = sum(r["amount"] for r in rows if r["type"] == "income")
    spent = sum(r["amount"] for r in rows if r["type"] == "expense")
    fee = spent * (o["fee_percent"] or 0) / 100
    by_cat = OrderedDict((c, 0.0) for c in EXPENSE_CATEGORIES)
    for r in rows:
        if r["type"] == "expense":
            by_cat[r["category"]] += r["amount"]
    c = o["currency"]
    by_cat = {k: _rnd(v, c) for k, v in by_cat.items() if v}
    balance = received - spent - fee
    return {
        "object": o["name"], "client": o["client"], "currency": c,
        "budget": _rnd(o["budget"] or 0, c), "fee_percent": o["fee_percent"] or 0,
        "received": _rnd(received, c), "spent": _rnd(spent, c), "fee": _rnd(fee, c),
        "balance": _rnd(balance, c),
        "budget_used_percent": round((spent + fee) / o["budget"] * 100, 1) if o["budget"] else None,
        "by_category": by_cat, "operations": len(rows),
        "period": {"from": date_from, "to": date_to},
    }


# ---------- отчёт клиенту ----------
def _period_text(t, date_from, date_to):
    f = lambda d: dt.date.fromisoformat(_parse_date(d)).strftime("%d.%m.%Y")  # noqa: E731
    if date_from and date_to:
        return f"{f(date_from)} – {f(date_to)}"
    if date_from:
        return f"{f(date_from)} – {dt.datetime.now(TZ):%d.%m.%Y}"
    if date_to:
        return f"… – {f(date_to)}"
    return t["all_time"]


def _safe(s):
    return re.sub(r"[^\w\-]+", "_", s, flags=re.UNICODE).strip("_")[:40] or "object"


def client_report(object, date_from=None, date_to=None, language="ru", formats=("docx", "xlsx"), company=""):
    t = L.get(language, L["ru"])
    o = find_object(object)
    s = summary(o["id"], date_from, date_to)
    rows = _tx(o["id"], date_from, date_to)
    incomes = [r for r in rows if r["type"] == "income"]
    expenses = [r for r in rows if r["type"] == "expense"]
    cur = o["currency"]
    period = _period_text(t, date_from, date_to)
    stamp = dt.datetime.now(TZ)
    base = OUT_DIR / f"report_{_safe(o['name'])}_{language}_{stamp:%Y%m%d_%H%M%S}"

    sum_rows = []
    if s["budget"]:
        sum_rows.append((t["budget"], s["budget"]))
    sum_rows += [(t["received"], s["received"]), (t["spent"], s["spent"])]
    if s["fee_percent"]:
        sum_rows.append((f"{t['fee']} ({s['fee_percent']:g}% {t['of_expenses']})", s["fee"]))
    if s["balance"] >= 0:
        sum_rows.append((t["balance"], s["balance"]))
    else:
        sum_rows.append((t["overspend"], -s["balance"]))

    files = []
    if "docx" in formats:
        files.append(_report_docx(base.with_suffix(".docx"), t, o, s, incomes, expenses, cur, period, stamp, sum_rows, company))
    if "xlsx" in formats:
        files.append(_report_xlsx(base.with_suffix(".xlsx"), t, o, s, incomes, expenses, cur, period, stamp, sum_rows))
    return {"summary": s, "files": files}


def _desc(r):
    d = r["description"] or ""
    if r.get("orig_currency"):
        op = "×" if r["orig_currency"] == "USD" else "÷"
        d += f" ({_money(r['orig_amount'], r['orig_currency'])} {r['orig_currency']} {op} {_money(r['rate'], 'USD').rstrip('0').rstrip('.')})"
    if r.get("receipt_file_id"):
        d += " 📎"
    return d


def _report_docx(path, t, o, s, incomes, expenses, cur, period, stamp, sum_rows, company):
    _m = lambda x: _money(x, cur)  # noqa: E731
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    for sec in d.sections:
        sec.left_margin, sec.right_margin = Cm(2), Cm(1.5)
        sec.top_margin = sec.bottom_margin = Cm(1.5)

    if company:
        p = d.add_paragraph(company)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.runs[0].italic = True
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(t["title"])
    r.bold, r.font.size = True, Pt(14)

    info = [(t["object"], o["name"]), (t["client"], o["client"] or "__________"),
            (t["address"], o["address"] or "—"), (t["period"], period),
            (t["report_date"], stamp.strftime("%d.%m.%Y"))]
    for k, v in info:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(f"{k}: ").bold = True
        p.add_run(str(v))

    def table(headers, data, bold_last=False, widths=None, font_pt=None):
        tb = d.add_table(rows=1, cols=len(headers))
        tb.style = "Table Grid"
        for i, hd in enumerate(headers):
            c = tb.rows[0].cells[i]
            c.text = hd
            c.paragraphs[0].runs[0].bold = True
        for ri, row in enumerate(data):
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                if bold_last and ri == len(data) - 1:
                    for run in cells[i].paragraphs[0].runs:
                        run.bold = True
        if widths:
            tb.autofit = False
            for i, col in enumerate(tb.columns):
                col.width = Cm(widths[i])
            for row in tb.rows:
                for i, c in enumerate(row.cells):
                    c.width = Cm(widths[i])
        if font_pt:
            for row in tb.rows:
                for c in row.cells:
                    for par in c.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(font_pt)
        for row in tb.rows:
            for i, c in enumerate(row.cells):
                if i > 0 and headers[i] in (f"{t['amount']}, {cur}", f"{t['price']}, {cur}", t["share"], t["qty"]):
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return tb

    def heading(text):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.add_run(text).bold = True

    heading(t["summary"])
    rows = [(k, _m(v)) for k, v in sum_rows]
    if s["budget_used_percent"] is not None:
        rows.append((t["budget_used"], f"{s['budget_used_percent']}%"))
    table(["", f"{t['amount']}, {cur}"], rows, widths=[11.5, 6])

    if s["by_category"]:
        heading(t["by_cat"])
        data = [(t["cats"][k], _m(v), f"{v / s['spent'] * 100:.1f}%") for k, v in s["by_category"].items()]
        data.append((t["total"], _m(s["spent"]), "100%"))
        table([t["category"], f"{t['amount']}, {cur}", t["share"]], data, bold_last=True, widths=[9, 6, 2.5])

    if incomes:
        heading(t["incomes"])
        data = [(i + 1, dt.date.fromisoformat(r["date"]).strftime("%d.%m.%Y"), t["cats"].get(r["category"], r["category"]),
                 _desc(r), _m(r["amount"])) for i, r in enumerate(incomes)]
        data.append(("", "", "", t["total"], _m(s["received"])))
        table([t["no"], t["date"], t["type"], t["desc"], f"{t['amount']}, {cur}"], data, bold_last=True,
              widths=[0.7, 2.3, 2.6, 8.4, 3.5])

    if expenses:
        heading(t["expenses"])
        data = []
        for i, r in enumerate(expenses):
            data.append((i + 1, dt.date.fromisoformat(r["date"]).strftime("%d.%m.%Y"),
                         t["cats"].get(r["category"], r["category"]), _desc(r),
                         f"{r['qty']:g}" if r["qty"] not in (None, "") else "", r["unit"] or "",
                         _m(r["price"]) if r["price"] not in (None, "") else "", _m(r["amount"])))
        data.append(("", "", "", t["total"], "", "", "", _m(s["spent"])))
        table([t["no"], t["date"], t["category"], t["desc"], t["qty"], t["unit"], f"{t['price']}, {cur}",
               f"{t['amount']}, {cur}"], data, bold_last=True, widths=[0.7, 2.1, 2.4, 5.0, 1.2, 1.2, 2.2, 2.7], font_pt=9.5)

    p = d.add_paragraph(t["note"])
    p.paragraph_format.space_before = Pt(10)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)
    d.add_paragraph()
    for who in (t["sign_c"], t["sign_z"]):
        d.add_paragraph(f"{who}: ____________________ / ____________________ /")
    d.save(path)
    return str(path)


def _report_xlsx(path, t, o, s, incomes, expenses, cur, period, stamp, sum_rows):
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0' if cur == "UZS" else '#,##0.00'

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font, cell.fill, cell.border = bold, head_fill, border
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    # Сводка
    ws = wb.active
    ws.title = t["sheet_sum"]
    ws["A1"] = t["title"]
    ws["A1"].font = Font(bold=True, size=13)
    info = [(t["object"], o["name"]), (t["client"], o["client"] or ""), (t["address"], o["address"] or ""),
            (t["period"], period), (t["report_date"], stamp.strftime("%d.%m.%Y"))]
    r = 3
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="")
    ws.cell(row=r, column=2, value=f"{t['amount']}, {cur}")
    style_header(ws, r, 2)
    for k, v in sum_rows:
        r += 1
        ws.cell(row=r, column=1, value=k).border = border
        c = ws.cell(row=r, column=2, value=v)
        c.number_format, c.border = money, border
    r += 2
    if s["by_category"]:
        ws.cell(row=r, column=1, value=t["category"])
        ws.cell(row=r, column=2, value=f"{t['amount']}, {cur}")
        ws.cell(row=r, column=3, value=t["share"])
        style_header(ws, r, 3)
        first = r + 1
        for k, v in s["by_category"].items():
            r += 1
            ws.cell(row=r, column=1, value=t["cats"][k]).border = border
            c = ws.cell(row=r, column=2, value=v)
            c.number_format, c.border = money, border
        last = r
        r += 1
        ws.cell(row=r, column=1, value=t["total"]).font = bold
        c = ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{last})")
        c.number_format, c.font = money, bold
        for rr in range(first, last + 1):
            c = ws.cell(row=rr, column=3, value=f"=B{rr}/B${r}")
            c.number_format, c.border = "0.0%", border
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 10

    # Поступления
    wi = wb.create_sheet(t["sheet_inc"])
    hdr = [t["no"], t["date"], t["type"], t["desc"], f"{t['amount']}, {cur}"]
    wi.append(hdr)
    style_header(wi, 1, len(hdr))
    for i, x in enumerate(incomes, 1):
        wi.append([i, dt.date.fromisoformat(x["date"]), t["cats"].get(x["category"], x["category"]), _desc(x), x["amount"]])
        wi.cell(row=i + 1, column=2).number_format = "DD.MM.YYYY"
        wi.cell(row=i + 1, column=5).number_format = money
    n = len(incomes) + 2
    wi.cell(row=n, column=4, value=t["total"]).font = bold
    c = wi.cell(row=n, column=5, value=f"=SUM(E2:E{max(n - 1, 2)})")
    c.number_format, c.font = money, bold
    for col, w in zip("ABCDE", (5, 12, 16, 40, 18)):
        wi.column_dimensions[col].width = w

    # Расходы
    we = wb.create_sheet(t["sheet_exp"])
    hdr = [t["no"], t["date"], t["category"], t["desc"], t["qty"], t["unit"], f"{t['price']}, {cur}", f"{t['amount']}, {cur}"]
    we.append(hdr)
    style_header(we, 1, len(hdr))
    for i, x in enumerate(expenses, 1):
        row = i + 1
        we.append([i, dt.date.fromisoformat(x["date"]), t["cats"].get(x["category"], x["category"]), _desc(x),
                   x["qty"], x["unit"], x["price"], None])
        we.cell(row=row, column=2).number_format = "DD.MM.YYYY"
        if x["qty"] not in (None, "") and x["price"] not in (None, ""):
            we.cell(row=row, column=8, value=f"=E{row}*G{row}")
        else:
            we.cell(row=row, column=8, value=x["amount"])
        we.cell(row=row, column=7).number_format = money
        we.cell(row=row, column=8).number_format = money
    n = len(expenses) + 2
    we.cell(row=n, column=4, value=t["total"]).font = bold
    c = we.cell(row=n, column=8, value=f"=SUM(H2:H{max(n - 1, 2)})")
    c.number_format, c.font = money, bold
    for col, w in zip("ABCDEFGH", (5, 12, 20, 40, 9, 7, 14, 16)):
        we.column_dimensions[col].width = w
    we.freeze_panes = "A2"
    wi.freeze_panes = "A2"

    wb.save(path)
    return str(path)
